#!/usr/bin/env python3
"""都道府県別エネルギー消費統計の 47 県の Excel から public/data/energy.json を書く。設計は docs/data-sources.md。"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/xlsx"
OUT = ROOT / "public/data/energy.json"

# 年度シートの列コード（1 行目）と行コード（A 列）。
CODE_ROW = 1
ENERGY_MARKER = "ｴﾈﾙｷﾞｰ単位表"
CARBON_MARKER = "炭素単位表"
TOTAL_ROW = 500000
NON_ENERGY_ROW = 950000
FINAL_COL = 900
ENERGY_USE_COL = 910
ELECTRICITY_LOSS_COL = 940
HEAT_LOSS_COL = 960
ALLOCATED_COL = 999
# 事業用水力・原子力は県へ配分できないので常に 0（Q&A F-1）。
ZERO_COLS = (550, 600)

# 指標シート: 1 人あたりブロックの人口の行と、最終エネルギー消費（電力・熱配分後）の GJ/人 の行。
INDICATOR_SHEET = "指標"
POPULATION_ROW = 12
PER_CAPITA_ROW = 14
FIRST_YEAR_COL = 7

TJ_PER_PJ = 1000


# JIS X 0401 の都道府県名。表のラベルは「東京」のような略称なので、正式名はここで持ち、略称が先頭に一致することを確かめる。
PREF_NAMES = [
    "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県", "茨城県", "栃木県", "群馬県",
    "埼玉県", "千葉県", "東京都", "神奈川県", "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県",
    "岐阜県", "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県",
    "鳥取県", "島根県", "岡山県", "広島県", "山口県", "徳島県", "香川県", "愛媛県", "高知県", "福岡県",
    "佐賀県", "長崎県", "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
]


@dataclass
class Item:
    code: str
    label: str
    keys: tuple[int, ...]


SECTORS = [
    Item("agriculture", "農林水産鉱建設業", (610000,)),
    Item("manufacturing", "製造業", (620000,)),
    Item("commercial", "業務他（第三次産業）", (650000,)),
    Item("household", "家庭", (700000,)),
    Item("car", "家庭の乗用車", (800000,)),
]

SOURCES = [
    Item("coal", "石炭・石炭製品", (100, 150)),
    Item("oil", "石油・石油製品", (200, 250)),
    Item("gas", "天然ガス", (400,)),
    Item("citygas", "都市ガス", (450,)),
    Item("renewable", "再生可能・未活用", (500,)),
    Item("electricity", "電力", (700,)),
    Item("heat", "熱", (800,)),
    Item("loss", "発電・熱供給の損失（配分）", (ELECTRICITY_LOSS_COL, HEAT_LOSS_COL)),
]

BASES = [("final", "最終エネルギー消費"), ("allocated", "電力・熱配分後")]

# 中分類を持つ部門。子は表の行順で、親の行の次から次の親の行の手前まで。
INDUSTRY_PARENTS = {610000: "agriculture", 620000: "manufacturing", 650000: "commercial"}
INDUSTRY_END = 700000


def clean(label: object) -> str:
    return str(label).strip().replace("\u3000", "").replace("･", "・")


def num(v: object) -> float:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    raise SystemExit(f"数値でない: {v!r}")


def close(a: float, b: float) -> bool:
    return abs(a - b) <= max(1e-6 * abs(b), 0.01)


class YearSheet:
    """年度シートのエネルギー単位表（TJ）。行コード × 列コードで引く。"""

    def __init__(self, ws: Worksheet) -> None:
        rows = list(ws.iter_rows(values_only=True))
        self.cols = {v: i for i, v in enumerate(rows[CODE_ROW - 1]) if isinstance(v, int)}
        markers = [i for i, r in enumerate(rows) if isinstance(r[3], str) and ENERGY_MARKER in r[3]]
        ends = [i for i, r in enumerate(rows) if isinstance(r[3], str) and CARBON_MARKER in r[3]]
        if len(markers) != 1 or len(ends) != 1:
            raise SystemExit(f"{ws.title}: エネルギー単位表の位置が分からない")
        self.rows = {r[0]: r for r in rows[markers[0] : ends[0]] if isinstance(r[0], int)}
        order = [c for c in self.rows if isinstance(c, int)]
        self.labels = {c: clean(self.rows[c][3]) for c in order}
        self.children: dict[int, list[int]] = {}
        parent: int | None = None
        for c in order:
            if c in INDUSTRY_PARENTS:
                parent = c
                self.children[c] = []
            elif c == INDUSTRY_END:
                parent = None
            elif parent is not None:
                self.children[parent].append(c)
        self.title = ws.title
        self.label = str(rows[0][0]).strip()

    def at(self, row: int, col: int) -> float:
        return num(self.rows[row][self.cols[col]])


def year_of(title: str) -> int:
    yy = int(title.removesuffix("FY"))
    return 1900 + yy if yy >= 90 else 2000 + yy


def industries(sheet: YearSheet) -> list[Item]:
    return [Item(str(c), sheet.labels[c], (c,)) for p in INDUSTRY_PARENTS for c in sheet.children[p]]


def read_pref(path: Path, industry: list[Item]) -> tuple[str, dict[int, dict], dict[int, float], int | None]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [YearSheet(wb[s]) for s in wb.sheetnames if s.endswith("FY")]
    label = str(next(wb[sheets[0].title].iter_rows(min_row=4, max_row=4, values_only=True))[0]).strip()
    label = label.replace("\u3000", "")

    ind = list(wb[INDICATOR_SHEET].iter_rows(max_row=PER_CAPITA_ROW, values_only=True))
    ind_years = [int(v) for v in ind[0][FIRST_YEAR_COL - 1 :] if isinstance(v, (int, float))]
    population = {y: num(ind[POPULATION_ROW - 1][FIRST_YEAR_COL - 1 + i]) for i, y in enumerate(ind_years)}
    per_capita = {y: num(ind[PER_CAPITA_ROW - 1][FIRST_YEAR_COL - 1 + i]) for i, y in enumerate(ind_years)}

    out: dict[int, dict] = {}
    provisional: int | None = None
    for s in sheets:
        year = year_of(s.title)
        if "暫定" in s.label:
            provisional = year
        if industries(s) != industry:
            raise SystemExit(f"{path.name} {s.title}: 中分類の並びが東京都と違う")
        for parent, children in s.children.items():
            for col in (FINAL_COL, ALLOCATED_COL):
                got = sum(s.at(c, col) for c in children)
                if not close(got, s.at(parent, col)):
                    raise SystemExit(f"{path.name} {s.title}: {parent} の中分類の和 {got} が合わない（列 {col}）")
        for col in ZERO_COLS:
            if s.at(TOTAL_ROW, col) != 0:
                raise SystemExit(f"{path.name} {s.title}: 列 {col} が 0 でない")

        final_total = s.at(TOTAL_ROW, FINAL_COL)
        allocated_total = s.at(TOTAL_ROW, ALLOCATED_COL)
        sector = {
            "final": [sum(s.at(r, FINAL_COL) for r in i.keys) for i in SECTORS],
            "allocated": [sum(s.at(r, ALLOCATED_COL) for r in i.keys) for i in SECTORS],
        }
        source_final = [0.0 if i.code == "loss" else sum(s.at(TOTAL_ROW, c) for c in i.keys) for i in SOURCES]
        # 配分後は非エネルギー利用を含まない（999 = 910 + 940 + 960）。燃料別の非エネルギー利用は 950000 行にある。
        source_allocated = [
            sum(s.at(TOTAL_ROW, c) for c in i.keys)
            - (0.0 if i.code == "loss" else sum(s.at(NON_ENERGY_ROW, c) for c in i.keys))
            for i in SOURCES
        ]

        checks = [
            ("部門の和（最終）", sum(sector["final"]), final_total),
            ("部門の和（配分後）", sum(sector["allocated"]), allocated_total),
            ("エネルギー種の和（最終）", sum(source_final), final_total),
            ("エネルギー種の和（配分後）", sum(source_allocated), allocated_total),
            (
                "配分後 = エネルギー利用 + 寄与損失",
                s.at(TOTAL_ROW, ENERGY_USE_COL) + s.at(TOTAL_ROW, ELECTRICITY_LOSS_COL) + s.at(TOTAL_ROW, HEAT_LOSS_COL),
                allocated_total,
            ),
        ]
        for name, got, want in checks:
            if not close(got, want):
                raise SystemExit(f"{path.name} {s.title}: {name} {got} が合計 {want} と合わない")
        official = per_capita[year]
        computed = allocated_total * 1e3 / population[year]
        if abs(computed - official) > 1e-6 * official:
            raise SystemExit(f"{path.name} {s.title}: 1人あたり {computed} が指標シート {official} と合わない")

        out[year] = {
            "sector": sector,
            "industry": {
                "final": [s.at(i.keys[0], FINAL_COL) for i in industry],
                "allocated": [s.at(i.keys[0], ALLOCATED_COL) for i in industry],
            },
            "source": {"final": source_final, "allocated": source_allocated},
        }
    return label, out, population, provisional


def cube(items: list[Item], key: str, areas: list[str], years: list[int], data: dict, parents: dict[str, str] | None = None) -> dict:
    values = [
        round(data[a][y][key][b][k] / TJ_PER_PJ, 3)
        for b, _ in BASES
        for a in areas
        for y in years
        for k in range(len(items))
    ]
    return {
        "items": [
            {"code": i.code, "label": i.label, "level": 1, **({"parent": parents[i.code]} if parents else {})}
            for i in items
        ],
        "dims": [
            {"name": "basis", "codes": [b for b, _ in BASES]},
            {"name": "area", "codes": areas},
            {"name": "year", "codes": [str(y) for y in years]},
            {"name": "item", "codes": [i.code for i in items]},
        ],
        "measures": {"value": values},
    }


def main() -> None:
    files = sorted(RAW.glob("*.xlsx"))
    if len(files) != 47:
        raise SystemExit(f"{RAW.relative_to(ROOT)} の xlsx が {len(files)} 本。npm run fetch を先に実行する。")

    tokyo = openpyxl.load_workbook(RAW / "13tokyo.xlsx", read_only=True, data_only=True)
    first = YearSheet(tokyo[next(n for n in tokyo.sheetnames if n.endswith("FY"))])
    industry = industries(first)
    parent_of = {str(c): INDUSTRY_PARENTS[p] for p, cs in first.children.items() for c in cs}

    areas: list[dict] = []
    data: dict[str, dict] = {}
    population: dict[str, dict[int, float]] = {}
    provisional: set[int | None] = set()
    for f in files:
        code = f.name[:2]
        short, years_data, pop, prov = read_pref(f, industry)
        label = PREF_NAMES[int(code) - 1]
        if not label.startswith(short):
            raise SystemExit(f"{f.name}: 表のラベル {short} が正式名 {label} と合わない")
        areas.append({"code": code, "label": label, "level": 1})
        data[code] = years_data
        population[code] = pop
        provisional.add(prov)
        print(f"{f.name}  {short} → {label}")

    years = sorted(next(iter(data.values())))
    if any(sorted(d) != years for d in data.values()):
        raise SystemExit("県によって年度がそろっていない")
    if len(provisional) != 1:
        raise SystemExit(f"暫定値の年度が県によって違う {provisional}")

    codes = [a["code"] for a in areas]
    if codes != [f"{i:02d}" for i in range(1, 48)]:
        raise SystemExit(f"県コードが 01–47 でない {codes}")

    # 47 都道府県計。全国の公表値ではなく、県の和（Q&A B-1: 総合エネルギー統計の全国値とは一致しない）。
    total = "00"
    data[total] = {
        y: {
            key: {b: [sum(data[c][y][key][b][k] for c in codes) for k in range(n)] for b, _ in BASES}
            for key, n in (("sector", len(SECTORS)), ("source", len(SOURCES)), ("industry", len(industry)))
        }
        for y in years
    }
    population[total] = {y: sum(population[c][y] for c in codes) for y in years}
    all_areas = [total, *codes]

    result = {
        "unit": "PJ",
        "provisional": provisional.pop(),
        "bases": [{"code": b, "label": l, "level": 0} for b, l in BASES],
        "areas": [{"code": total, "label": "47都道府県計", "level": 0}, *areas],
        "population": {
            "dims": [
                {"name": "area", "codes": all_areas},
                {"name": "year", "codes": [str(y) for y in years]},
            ],
            "measures": {"value": [round(population[a][y]) for a in all_areas for y in years]},
        },
        "sector": cube(SECTORS, "sector", all_areas, years, data),
        "source": cube(SOURCES, "source", all_areas, years, data),
        "industry": cube(industry, "industry", all_areas, years, data, parent_of),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"energy.json  {len(all_areas)} areas × {len(years)} years ({years[0]}–{years[-1]}, 暫定 {result['provisional']})")


if __name__ == "__main__":
    main()
